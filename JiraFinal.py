import requests
from docx import Document
from requests.auth import HTTPBasicAuth
from dotenv import load_dotenv
import os
from datetime import datetime

# === LOAD ENV VARIABLES ===
load_dotenv()

JIRA_BASE_URL = os.getenv("JIRA_BASE_URL")
JIRA_EMAIL = os.getenv("JIRA_EMAIL")
JIRA_API_TOKEN = os.getenv("JIRA_API_TOKEN")
MODIFICATION_FIELD_ID = os.getenv("MODIFICATION_FIELD_ID")
OUTPUT_PATH = os.getenv("OUTPUT_PATH", "C:/Users/AP109100/Desktop/pytest/JiraNotesSummary.docx")

# === JQL QUERY ===
JQL_QUERY = 'fixVersion = 64766 AND project = OLK ORDER BY created ASC'

# === AUTHENTICATION ===
auth = HTTPBasicAuth(JIRA_EMAIL, JIRA_API_TOKEN)
headers = {
    "Accept": "application/json"
}


# === FETCH ISSUES ===
def fetch_issues():
    url = f"{JIRA_BASE_URL}/rest/api/3/search"
    params = {
        "jql": JQL_QUERY,
        "fields": MODIFICATION_FIELD_ID,
        "maxResults": 100
    }
    try:
        response = requests.get(url, headers=headers, auth=auth, params=params, timeout=10)
        response.raise_for_status()
        return response.json().get("issues", [])
    except requests.exceptions.RequestException as err:
        print(f"Network error: {err}")
        return []


# === GROUP ISSUES BY HEADING ===
def group_issues_by_heading(issuesdoc, field_id):
    grouped = {}

    for issue in issuesdoc:
        issue_key = issue["key"]
        field_data = issue["fields"].get(field_id)

        if not field_data:
            continue

        for block in field_data.get("content", []):
            if block["type"] == "paragraph":
                lines = []
                line = ""
                for inline in block.get("content", []):
                    if inline["type"] == "text":
                        line += inline["text"]
                    elif inline["type"] == "hardBreak":
                        lines.append(line.strip())
                        line = ""
                if line:
                    lines.append(line.strip())

                if not lines:
                    continue

                heading = lines[0]  # first line is heading
                details = lines[1:] if len(lines) > 1 else ["No details provided"]

                if heading not in grouped:
                    grouped[heading] = []

                for detail in details:
                    grouped[heading].append((detail, issue_key))

    return grouped


# === CREATE WORD DOCUMENT ===
def create_word_doc(issuesdoc):
    doc = Document()
    doc.add_heading("Jira Modification Summary", level=1)

    # Group issues by heading
    grouped_data = group_issues_by_heading(issuesdoc, MODIFICATION_FIELD_ID)

    for heading, items in grouped_data.items():
        # Add heading (bold + underline)
        p = doc.add_paragraph()
        run = p.add_run(heading)
        run.bold = True
        run.underline = True

        # Add bullet points for all details
        for detail, issue_key in items:
            bullet = doc.add_paragraph(style="List Bullet")
            bullet.add_run(f"{detail} ({issue_key})")

    # === Add timestamp to filename ===
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    base, ext = os.path.splitext(OUTPUT_PATH)
    output_with_timestamp = f"{base}_{timestamp}{ext}"

    try:
        doc.save(output_with_timestamp)
        return output_with_timestamp
    except Exception as error:
        print(f"File save error: {error}")
        return None


# === MAIN EXECUTION ===
if __name__ == "__main__":
    try:
        issues = fetch_issues()
        output_file = create_word_doc(issues)
        print(f" Document created: {output_file}")
    except Exception as e:
        print(f" Error: {e}")
