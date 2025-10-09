import requests
from docx import Document
from requests.auth import HTTPBasicAuth
from dotenv import load_dotenv
import os

# === LOAD ENV VARIABLES ===
load_dotenv()

JIRA_BASE_URL = os.getenv("JIRA_BASE_URL")
JIRA_EMAIL = os.getenv("JIRA_EMAIL")
JIRA_API_TOKEN = os.getenv("JIRA_API_TOKEN")
MODIFICATION_FIELD_ID = os.getenv("MODIFICATION_FIELD_ID")  # Use env variable name, not field label

# === JQL QUERY ===
JQL_QUERY = 'fixversion = "15.4 MediaWindow Editor" AND project = OLK AND type NOT IN (Bug, Task) AND labels NOT IN (Automation_TechDebt, Automation, DevOnly, QPlan, DevOps, Documentation, TestAutomationDebt, Sapient) ORDER BY created ASC'

# === AUTHENTICATION ===
auth = HTTPBasicAuth(JIRA_EMAIL, JIRA_API_TOKEN)
headers = {
    "Accept": "application/json"
}

# === FETCH ISSUES ===
def fetch_issues():
    url = f"{JIRA_BASE_URL}/rest/api/3/search"
    params = {
        "jql": JQL_QUERY,
        "fields": [MODIFICATION_FIELD_ID],
        "maxResults": 100
    }
    response = requests.get(url, headers=headers, auth=auth, params=params)
    response.raise_for_status()
    return response.json()["issues"]

# === CREATE WORD DOCUMENT ===
def create_word_doc(issues):
    doc = Document()
    doc.add_heading("Jira Modification Summary", level=1)

    for issue in issues:
        key = issue["key"]
        modification_info = issue["fields"].get(MODIFICATION_FIELD_ID, "No info provided")

        doc.add_heading(f"{key}", level=2)
        doc.add_paragraph(f"{modification_info}")

    doc.save("C:/Users/AP109100/Desktop/pytest/JiraNotesSummary.docx")

# === MAIN EXECUTION ===
if __name__ == "__main__":
    try:
        issues = fetch_issues()
        create_word_doc(issues)
        print("Document created: JiraNotesSummary.docx")
    except Exception as e:
        print(f"Error: {e}")
